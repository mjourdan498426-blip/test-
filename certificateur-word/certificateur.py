"""Certificateur de documents Word.

Certifier : ajoute un QR code à la fin d'une copie du Word, la convertit en PDF
(LibreOffice), enregistre l'empreinte SHA-256 du PDF dans registre.json et
dépose le PDF dans un dossier « Certifies » à côté du Word d'origine.

Vérifier : recalcule l'empreinte d'un PDF et la compare au registre.
"""

import datetime
import hashlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import uuid
import zipfile
from pathlib import Path

DOSSIER_APP = Path(__file__).resolve().parent
REGISTRE = DOSSIER_APP / "registre.json"
NOM_DOSSIER_SORTIE = "Certifies"
PREFIXE_QR = "DOCQR"
LIEN_LIBREOFFICE = "https://fr.libreoffice.org/download/telecharger-libreoffice/"


class ErreurUtilisateur(Exception):
    """Erreur à afficher telle quelle à l'utilisateur."""


# --------------------------------------------------------------------------
# Outils
# --------------------------------------------------------------------------

def sha256_fichier(chemin):
    h = hashlib.sha256()
    with open(chemin, "rb") as f:
        for bloc in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloc)
    return h.hexdigest()


def trouver_libreoffice():
    """Renvoie le chemin de l'exécutable LibreOffice, ou None s'il est absent."""
    for nom in ("soffice", "libreoffice"):
        chemin = shutil.which(nom)
        if chemin:
            return chemin
    candidats = [
        Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "LibreOffice/program/soffice.exe",
        Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")) / "LibreOffice/program/soffice.exe",
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    for c in candidats:
        if c.is_file():
            return str(c)
    return None


def message_libreoffice_absent():
    return (
        "LibreOffice n'est pas installé sur cet ordinateur.\n\n"
        "Il est nécessaire pour fabriquer les PDF certifiés.\n\n"
        "Pour l'installer :\n"
        f"1. Ouvrez la page {LIEN_LIBREOFFICE}\n"
        "2. Cliquez sur « Télécharger », puis ouvrez le fichier téléchargé.\n"
        "3. Suivez l'installation en gardant les choix proposés.\n"
        "4. Relancez cette application.\n\n"
        "En attendant, la vérification des PDF reste possible."
    )


def lire_registre():
    if not REGISTRE.exists():
        return {"documents": []}
    try:
        with open(REGISTRE, encoding="utf-8") as f:
            donnees = json.load(f)
        if not isinstance(donnees.get("documents"), list):
            raise ValueError
        return donnees
    except (ValueError, OSError, AttributeError):
        raise ErreurUtilisateur(
            f"Le registre ({REGISTRE.name}) est illisible ou abîmé.\n"
            "Rien n'a été modifié. Restaurez une copie de sauvegarde de ce fichier."
        )


def ecrire_registre(donnees):
    # Écriture dans un fichier temporaire puis remplacement : le registre
    # n'est jamais laissé à moitié écrit, même en cas de coupure.
    tmp = REGISTRE.with_suffix(".json.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(donnees, f, ensure_ascii=False, indent=2)
    os.replace(tmp, REGISTRE)


def verifier_docx(chemin):
    """Lève ErreurUtilisateur si le fichier n'est pas un .docx valide."""
    chemin = Path(chemin)
    if not chemin.is_file():
        raise ErreurUtilisateur(f"Fichier introuvable :\n{chemin}")
    if chemin.suffix.lower() != ".docx":
        raise ErreurUtilisateur(
            f"« {chemin.name} » n'est pas un document Word (.docx).\n\n"
            "Seuls les fichiers .docx peuvent être certifiés. Si c'est un ancien "
            "fichier .doc, ouvrez-le dans Word puis « Enregistrer sous » au format .docx."
        )
    try:
        with zipfile.ZipFile(chemin) as z:
            if "word/document.xml" not in z.namelist():
                raise zipfile.BadZipFile
        import docx
        docx.Document(str(chemin))
    except Exception:
        raise ErreurUtilisateur(
            f"« {chemin.name} » porte l'extension .docx mais n'est pas un document "
            "Word valide (fichier abîmé ou d'un autre type)."
        )


def chemin_libre(dossier, nom_base, extension):
    chemin = dossier / f"{nom_base}{extension}"
    n = 2
    while chemin.exists():
        chemin = dossier / f"{nom_base}_{n}{extension}"
        n += 1
    return chemin


def convertir_en_pdf(soffice, docx_path, dossier_sortie):
    # Profil LibreOffice temporaire : la conversion fonctionne même si
    # LibreOffice est déjà ouvert par ailleurs.
    profil = Path(dossier_sortie) / "profil_lo"
    commande = [
        soffice,
        f"-env:UserInstallation={profil.as_uri()}",
        "--headless", "--norestore",
        "--convert-to", "pdf",
        "--outdir", str(dossier_sortie),
        str(docx_path),
    ]
    options = {}
    if sys.platform == "win32":
        options["creationflags"] = subprocess.CREATE_NO_WINDOW
    try:
        subprocess.run(commande, capture_output=True, timeout=300, check=False, **options)
    except subprocess.TimeoutExpired:
        raise ErreurUtilisateur("La conversion en PDF a pris trop de temps et a été interrompue.")
    pdf = Path(dossier_sortie) / (Path(docx_path).stem + ".pdf")
    if not pdf.is_file():
        raise ErreurUtilisateur(
            "LibreOffice n'a pas réussi à convertir le document en PDF.\n"
            "Vérifiez que le document s'ouvre correctement dans Word ou LibreOffice."
        )
    return pdf


# --------------------------------------------------------------------------
# Certifier / Vérifier
# --------------------------------------------------------------------------

def certifier(chemin_word):
    """Certifie un .docx. Renvoie l'entrée ajoutée au registre et le chemin du PDF."""
    import docx
    import qrcode
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    chemin_word = Path(chemin_word).resolve()
    verifier_docx(chemin_word)
    soffice = trouver_libreoffice()
    if not soffice:
        raise ErreurUtilisateur(message_libreoffice_absent())
    registre = lire_registre()  # échoue tôt si le registre est abîmé

    reference = str(uuid.uuid4())
    maintenant = datetime.datetime.now()
    empreinte_word = sha256_fichier(chemin_word)
    contenu_qr = f"{PREFIXE_QR}|{reference}|{empreinte_word[:16]}"

    with tempfile.TemporaryDirectory(prefix="certif_") as tmp:
        tmp = Path(tmp)
        # On travaille sur une copie : l'original n'est jamais ouvert en écriture.
        copie = tmp / chemin_word.name
        shutil.copyfile(chemin_word, copie)

        document = docx.Document(str(copie))

        image_qr = io.BytesIO()
        qrcode.make(contenu_qr, border=2).save(image_qr, format="PNG")
        image_qr.seek(0)

        p_qr = document.add_paragraph()
        p_qr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_qr.paragraph_format.space_before = Pt(18)
        p_qr.paragraph_format.keep_with_next = True  # QR et légende sur la même page
        p_qr.add_run().add_picture(image_qr, width=Cm(3))

        p_legende = document.add_paragraph()
        run = p_legende.add_run(
            f"Document certifié le {maintenant:%d/%m/%Y} - Réf. {reference[:8]}"
        )
        run.font.size = Pt(8)
        run.italic = True

        document.save(str(copie))
        pdf_tmp = convertir_en_pdf(soffice, copie, tmp)

        dossier_sortie = chemin_word.parent / NOM_DOSSIER_SORTIE
        dossier_sortie.mkdir(exist_ok=True)
        pdf_final = chemin_libre(dossier_sortie, f"{chemin_word.stem}_certifie", ".pdf")
        shutil.move(str(pdf_tmp), pdf_final)

    entree = {
        "reference": reference,
        "fichier": chemin_word.name,
        "fichier_pdf": pdf_final.name,
        "empreinte_sha256": sha256_fichier(pdf_final),
        "empreinte_word_sha256": empreinte_word,
        "date": maintenant.isoformat(timespec="seconds"),
    }
    registre["documents"].append(entree)
    ecrire_registre(registre)
    return entree, pdf_final


def verifier(chemin_pdf):
    """Renvoie l'entrée du registre si le PDF est authentique, sinon None."""
    chemin_pdf = Path(chemin_pdf)
    if not chemin_pdf.is_file():
        raise ErreurUtilisateur(f"Fichier introuvable :\n{chemin_pdf}")
    if chemin_pdf.suffix.lower() != ".pdf":
        raise ErreurUtilisateur(
            f"« {chemin_pdf.name} » n'est pas un PDF.\n"
            "Pour vérifier un document, choisissez le PDF certifié (dossier « Certifies »)."
        )
    empreinte = sha256_fichier(chemin_pdf)
    for entree in lire_registre()["documents"]:
        if entree.get("empreinte_sha256") == empreinte:
            return entree
    return None


def date_lisible(iso):
    try:
        return datetime.datetime.fromisoformat(iso).strftime("%d/%m/%Y à %Hh%M")
    except (TypeError, ValueError):
        return str(iso)


# --------------------------------------------------------------------------
# Interface
# --------------------------------------------------------------------------

def lancer_interface(fichiers_au_demarrage):
    import tkinter as tk
    from tkinter import filedialog, messagebox

    racine = tk.Tk()
    # Glisser-déposer facultatif : si le module manque ou ne se charge pas
    # sur ce Mac, la fenêtre fonctionne quand même avec les boutons.
    try:
        from tkinterdnd2 import DND_FILES, TkinterDnD
        TkinterDnD.require(racine)
        glisser_ok = True
    except Exception:
        glisser_ok = False

    racine.title("Certificateur de documents Word")
    racine.geometry("620x560")
    racine.minsize(520, 480)
    racine.configure(bg="#f4f5f7")

    police_titre = ("Helvetica", 15, "bold")
    police = ("Helvetica", 11)

    tk.Label(racine, text="Certificateur de documents", font=("Helvetica", 18, "bold"),
             bg="#f4f5f7", fg="#1f2937").pack(pady=(16, 4))

    soffice = trouver_libreoffice()
    occupe = {"valeur": False}

    cadre_zones = tk.Frame(racine, bg="#f4f5f7")
    cadre_zones.pack(fill="x", padx=16, pady=8)
    cadre_zones.columnconfigure(0, weight=1, uniform="z")
    cadre_zones.columnconfigure(1, weight=1, uniform="z")

    resultat = tk.Label(racine, text="", font=("Helvetica", 12), wraplength=560,
                        justify="center", bg="#f4f5f7", padx=12, pady=14)
    resultat.pack(fill="both", expand=True, padx=16, pady=(4, 16))

    def afficher(texte, couleur_fond="#f4f5f7", couleur_texte="#1f2937"):
        resultat.configure(text=texte, bg=couleur_fond, fg=couleur_texte)

    def executer_en_fond(travail, quand_fini):
        """Lance un traitement sans figer la fenêtre."""
        if occupe["valeur"]:
            messagebox.showinfo("Patientez", "Un traitement est déjà en cours.")
            return
        occupe["valeur"] = True
        racine.configure(cursor="watch")

        def fil():
            try:
                res, err = travail(), None
            except ErreurUtilisateur as e:
                res, err = None, str(e)
            except Exception as e:  # erreur imprévue : on l'affiche quand même
                res, err = None, f"Erreur inattendue : {e}"
            racine.after(0, lambda: fin(res, err))

        def fin(res, err):
            occupe["valeur"] = False
            racine.configure(cursor="")
            if err:
                afficher("❌ " + err.split("\n")[0], "#fde8e8", "#9b1c1c")
                messagebox.showerror("Impossible", err)
            else:
                quand_fini(res)

        threading.Thread(target=fil, daemon=True).start()

    def lancer_certification(chemin):
        if not soffice:
            messagebox.showwarning("LibreOffice manquant", message_libreoffice_absent())
            return
        afficher(f"⏳ Certification de « {Path(chemin).name} » en cours…\n"
                 "(cela peut prendre quelques secondes)")

        def fini(res):
            entree, pdf = res
            afficher(
                "✅ Document certifié\n\n"
                f"Fichier : {entree['fichier']}\n"
                f"Référence : {entree['reference'][:8]}\n\n"
                f"PDF certifié enregistré dans le dossier « {NOM_DOSSIER_SORTIE} » :\n{pdf}",
                "#e6f6ec", "#14532d")

        executer_en_fond(lambda: certifier(chemin), fini)

    def lancer_verification(chemin):
        afficher(f"⏳ Vérification de « {Path(chemin).name} »…")

        def fini(entree):
            if entree:
                afficher(
                    "✅ Document authentique et non modifié\n\n"
                    f"Document d'origine : {entree['fichier']}\n"
                    f"Certifié le {date_lisible(entree['date'])}\n"
                    f"Référence : {entree['reference'][:8]}",
                    "#e6f6ec", "#14532d")
            else:
                afficher(
                    "❌ Document modifié ou inconnu\n\n"
                    f"« {Path(chemin).name} » ne correspond à aucun document certifié.\n"
                    "Il a pu être modifié, même légèrement, après sa certification.",
                    "#fde8e8", "#9b1c1c")

        executer_en_fond(lambda: verifier(chemin), fini)

    def creer_zone(colonne, titre, explication, types, action):
        zone = tk.Frame(cadre_zones, bg="white", highlightthickness=2,
                        highlightbackground="#cbd5e1", padx=12, pady=16)
        zone.grid(row=0, column=colonne, sticky="nsew", padx=6)
        tk.Label(zone, text=titre, font=police_titre, bg="white", fg="#1f2937").pack()
        tk.Label(zone, text=explication, font=police, bg="white", fg="#475569",
                 wraplength=240).pack(pady=(6, 10))

        def choisir():
            chemin = filedialog.askopenfilename(title=titre, filetypes=types)
            if chemin:
                action(chemin)

        tk.Button(zone, text="Choisir un fichier", font=police, command=choisir,
                  padx=10, pady=4).pack()
        if glisser_ok:
            tk.Label(zone, text="ou glissez le fichier ici", font=("Helvetica", 10, "italic"),
                     bg="white", fg="#64748b").pack(pady=(10, 0))

            def deposer(event):
                fichiers = racine.tk.splitlist(event.data)
                if fichiers:
                    action(fichiers[0])

            for w in [zone] + list(zone.winfo_children()):
                w.drop_target_register(DND_FILES)
                w.dnd_bind("<<Drop>>", deposer)
        return zone

    creer_zone(0, "Certifier un Word",
               "Ajoute un QR code et crée un PDF certifié.",
               [("Document Word", "*.docx")], lancer_certification)
    creer_zone(1, "Vérifier un PDF",
               "Contrôle qu'un PDF certifié n'a pas été modifié.",
               [("Document PDF", "*.pdf")], lancer_verification)

    if not soffice:
        afficher("⚠️ LibreOffice n'est pas installé : la certification est indisponible.\n"
                 "La vérification des PDF fonctionne.", "#fff7e0", "#7c4a03")
        racine.after(300, lambda: messagebox.showwarning("LibreOffice manquant",
                                                         message_libreoffice_absent()))
    elif not glisser_ok:
        afficher("Cliquez sur « Choisir un fichier » pour commencer.")
    else:
        afficher("Glissez un fichier dans l'une des deux zones,\n"
                 "ou cliquez sur « Choisir un fichier ».")

    # Fichier glissé directement sur le lanceur : on choisit l'action selon le type.
    def traiter_demarrage():
        for f in fichiers_au_demarrage[:1]:
            if f.lower().endswith(".pdf"):
                lancer_verification(f)
            else:
                lancer_certification(f)

    racine.after(400, traiter_demarrage)
    racine.mainloop()


if __name__ == "__main__":
    lancer_interface(sys.argv[1:])
